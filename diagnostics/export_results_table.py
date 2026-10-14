"""Export rollout results to a spreadsheet in the team's results_raw format.

Sheet "results_raw": one row per episode
  idx | type | task | instruction | swap | left | right | answer_side | success | robot_went (extra)
Sheet "summary": success rate per type x task, split by swap / no swap / total (live formulas).

Task names follow the team convention: knowledge -> "Semantic reasoning",
tile_object / explicit_object -> "Explicit object", explicit_spatial -> "Explicit spatial",
neutral -> "Neutral (control)".

Run from the repo root:  python -m diagnostics.export_results_table --out outputs/results_raw.xlsx
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from diagnostics.intent_classifier import classify_intent

CARROT = Path(__file__).resolve().parents[1] / "ManiSkill" / "mani_skill" / "assets" / "carrot"
# asset -> "type" label in the sheet
ASSETS = {
    "photo_state_v1": "state",
    "photo_attribute_v1": "attribute",
    "photo_emotion_v1": "emotion",
    "semantic_repeated_v2": "drawings pilot",
    "emotion_pilot_v1": "emotion (drawn)",
}
TASKS = {"paper_style": "Paper-style object", "tile_object": "Explicit object", "explicit_object": "Explicit object",
         "explicit_spatial": "Explicit spatial", "knowledge": "Semantic reasoning",
         "neutral": "Neutral (control)"}
TASK_ORDER = ["Paper-style object", "Explicit object", "Explicit spatial", "Semantic reasoning", "Neutral (control)"]
FONT = "Arial"


def load_rows(outputs: Path) -> list[dict]:
    rows = []
    for asset, type_label in ASSETS.items():
        pairs_path = CARROT / asset / "pairs.json"
        if not pairs_path.exists():
            continue
        pairs = {int(p["index"]): p for p in json.loads(pairs_path.read_text())}
        seen = {}
        for log in sorted(outputs.glob(f"openvla-{asset}-*/diagnostics/episode_log.jsonl")):
            for line in log.read_text().splitlines():
                if line.strip():
                    r = json.loads(line)
                    seen[(r["condition"], int(r["question_id"]), r["layout"])] = r
        for (cond, qid, layout), r in seen.items():
            p = pairs[qid]
            swap = int(layout == "swap")
            left, right = (p["right"], p["left"]) if swap else (p["left"], p["right"])
            went, _ = classify_intent(r)
            rows.append({"idx": qid, "type": type_label, "task": TASKS[cond], "instruction": r["instruction"],
                         "swap": swap, "left": left, "right": right, "answer_side": r["correct_side"],
                         "success": int(bool(r["task_success"])), "robot_went": went})
    type_order = list(ASSETS.values())
    rows.sort(key=lambda x: (type_order.index(x["type"]), TASK_ORDER.index(x["task"]), x["swap"], x["idx"]))
    return rows


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--outputs-root", default="outputs")
    ap.add_argument("--out", default="outputs/results_raw.xlsx")
    ap.add_argument("--assets", nargs="+", default=None, help="only these test assets, e.g. photo_state_v1")
    ap.add_argument("--no-neutral", action="store_true", help="leave out the neutral control")
    args = ap.parse_args()
    if args.assets:
        for a in list(ASSETS):
            if a not in args.assets:
                del ASSETS[a]
    rows = load_rows(Path(args.outputs_root))
    if args.no_neutral:
        rows = [r for r in rows if r["task"] != "Neutral (control)"]

    wb = Workbook()
    raw = wb.active
    raw.title = "results_raw"
    cols = ["idx", "type", "task", "instruction", "swap", "left", "right", "answer_side", "success", "robot_went"]
    raw.append(cols)
    for r in rows:
        raw.append([r[c] for c in cols])
    n = len(rows) + 1
    header = Font(name=FONT, bold=True)
    for cell in raw[1]:
        cell.font = header
    for row in raw.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
    widths = {"A": 6, "B": 18, "C": 20, "D": 80, "E": 6, "F": 24, "G": 24, "H": 12, "I": 9, "J": 12}
    for col, w in widths.items():
        raw.column_dimensions[col].width = w
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = f"A1:J{n}"
    green = PatternFill("solid", start_color="D9EAD3")
    red = PatternFill("solid", start_color="F4CCCC")
    raw.conditional_formatting.add(f"I2:I{n}", CellIsRule(operator="equal", formula=["1"], fill=green))
    raw.conditional_formatting.add(f"I2:I{n}", CellIsRule(operator="equal", formula=["0"], fill=red))
    raw.cell(row=1, column=12, value="robot_went = side the arm actually moved toward (from the trajectory); "
                                     "success = cube placed on the correct tile. One run per episode (seed 0).").font = Font(name=FONT, italic=True)

    summ = wb.create_sheet("summary")
    summ.append(["type", "task", "swap", "no swap", "total", "episodes"])
    for cell in summ[1]:
        cell.font = header
    combos = []
    for r in rows:
        if (r["type"], r["task"]) not in combos:
            combos.append((r["type"], r["task"]))
    rng = lambda c: f"results_raw!${c}$2:${c}${n}"
    for i, (t, task) in enumerate(combos, start=2):
        summ.cell(row=i, column=1, value=t)
        summ.cell(row=i, column=2, value=task)
        for col, swap in ((3, 1), (4, 0)):
            summ.cell(row=i, column=col, value=(
                f'=IFERROR(AVERAGEIFS({rng("I")},{rng("B")},$A{i},{rng("C")},$B{i},{rng("E")},{swap}),"")'))
        summ.cell(row=i, column=5, value=f'=IFERROR(AVERAGEIFS({rng("I")},{rng("B")},$A{i},{rng("C")},$B{i}),"")')
        summ.cell(row=i, column=6, value=f'=COUNTIFS({rng("B")},$A{i},{rng("C")},$B{i})')
        for col in (3, 4, 5):
            summ.cell(row=i, column=col).number_format = "0.0%"
    for row in summ.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.alignment = Alignment(horizontal="left" if cell.column <= 2 else "right")
    for col, w in zip("ABCDEF", (18, 20, 9, 9, 9, 10)):
        summ.column_dimensions[col].width = w
    note_row = len(combos) + 3
    summ.cell(row=note_row, column=1, value="Success rate = mean of results_raw!success for that type, task and swap value "
                                            "(formulas update if results_raw changes).").font = Font(name=FONT, italic=True)

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"{len(rows)} episodes, {len(combos)} type x task rows -> {args.out}")


if __name__ == "__main__":
    main()
